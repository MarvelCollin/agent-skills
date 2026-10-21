import argparse
import datetime
import getpass
import importlib.util
import io
import json
import os
import shutil
import sys
import zipfile
from pathlib import Path

LABELS = {
    "id": {
        "asked": "Masalah",
        "changed": "Diubah",
        "checked": "Dicek",
        "before": "Sebelum",
        "after": "Sesudah",
        "no_image": "Tidak ada gambar.",
    },
    "en": {
        "asked": "Asked",
        "changed": "Changed",
        "checked": "Checked",
        "before": "Before",
        "after": "After",
        "no_image": "No image.",
    },
}

BOX_COLOR = (220, 38, 38)
THUMB_SIZE = (640, 452)


def fail(message, code=2):
    print(f"ERROR: {message}")
    sys.exit(code)


def display_name():
    override = os.environ.get("REVISION_DOC_AUTHOR", "").strip()
    if override:
        return override
    if sys.platform == "win32":
        try:
            import ctypes

            size = ctypes.c_ulong(256)
            buffer = ctypes.create_unicode_buffer(256)
            if ctypes.windll.secur32.GetUserNameExW(3, buffer, ctypes.byref(size)) and buffer.value.strip():
                return buffer.value.strip()
        except Exception:
            pass
    else:
        try:
            import pwd

            gecos = pwd.getpwnam(getpass.getuser()).pw_gecos.split(",")[0].strip()
            if gecos:
                return gecos
        except Exception:
            pass
    return getpass.getuser()


def load_json(path, required=True):
    if not path.exists():
        if required:
            fail(f"{path} not found")
        return {}
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as error:
        fail(f"{path} is not valid JSON: {error}")


def validate(doc):
    problems = []
    for key in ("title", "items"):
        if key not in doc:
            problems.append(f"items.json needs '{key}'")
    for index, item in enumerate(doc.get("items", []), start=1):
        if not item.get("title"):
            problems.append(f"item {index} needs a title")
        for key in ("asked", "changed", "checked"):
            if not isinstance(item.get(key, []), list):
                problems.append(f"item {index} '{key}' must be a list")
        for side in ("before", "after"):
            if not item.get(side) and not item.get(f"{side}_note"):
                problems.append(f"item {index} needs '{side}' or '{side}_note'")
    if problems:
        for problem in problems:
            print(f"INVALID: {problem}")
        sys.exit(2)


def boxed_image(root, side, relative, boxes, viewport):
    from PIL import Image, ImageDraw

    source = root / side / relative
    if not source.exists():
        print(f"MISSING_IMAGE: {side}/{relative}")
        return None
    image = Image.open(source).convert("RGB")
    scale = image.width / viewport
    draw = ImageDraw.Draw(image)
    width = max(3, round(4 * scale))
    key = relative.replace("\\", "/")
    for box in boxes.get(key, []):
        x0 = box["x"] * scale
        y0 = box["y"] * scale
        x1 = x0 + box["w"] * scale
        y1 = y0 + box["h"] * scale
        draw.rectangle([x0 - width, y0 - width, x1 + width, y1 + width], outline=BOX_COLOR, width=width)
    target = root / "boxed" / side / key
    target.parent.mkdir(parents=True, exist_ok=True)
    image.save(target)
    return target


def font(size, bold=False):
    from PIL import ImageFont

    names = ["arialbd.ttf", "Arial Bold.ttf", "DejaVuSans-Bold.ttf"] if bold else ["arial.ttf", "Arial.ttf", "DejaVuSans.ttf"]
    for name in names:
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            continue
    try:
        return ImageFont.load_default(size=size)
    except TypeError:
        return ImageFont.load_default()


def wrap(draw, text, face, width):
    lines, line = [], ""
    for word in text.split():
        candidate = f"{line} {word}".strip()
        if draw.textlength(candidate, font=face) <= width or not line:
            line = candidate
        else:
            lines.append(line)
            line = word
    if line:
        lines.append(line)
    return lines


def thumbnail(doc, first_pair, labels):
    from PIL import Image, ImageDraw

    page = Image.new("RGB", THUMB_SIZE, "white")
    draw = ImageDraw.Draw(page)
    margin = 28
    width = THUMB_SIZE[0] - margin * 2
    y = margin
    for line in wrap(draw, doc["title"], font(26, bold=True), width)[:2]:
        draw.text((margin, y), line, fill=(17, 24, 39), font=font(26, bold=True))
        y += 32
    for line in wrap(draw, doc.get("intro", ""), font(12), width)[:3]:
        draw.text((margin, y + 4), line, fill=(75, 85, 99), font=font(12))
        y += 16
    y += 16
    column = (width - 16) // 2
    for index, (label, path) in enumerate(zip((labels["before"], labels["after"]), first_pair)):
        x = margin + index * (column + 16)
        draw.text((x, y), label, fill=(17, 24, 39), font=font(13, bold=True))
        frame_top = y + 20
        frame_height = THUMB_SIZE[1] - frame_top - margin
        if path:
            shot = Image.open(path).convert("RGB")
            shot.thumbnail((column, frame_height))
            page.paste(shot, (x, frame_top))
            draw.rectangle([x, frame_top, x + shot.width - 1, frame_top + shot.height - 1], outline=(209, 213, 219), width=1)
        else:
            draw.rectangle([x, frame_top, x + column, frame_top + frame_height], outline=(209, 213, 219), width=1)
    buffer = io.BytesIO()
    page.save(buffer, "JPEG", quality=88)
    return buffer.getvalue()


def embed_thumbnail(document, blob):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part

    package = document.part.package
    for rel in list(package.rels.values()):
        if rel.reltype == RT.THUMBNAIL:
            part = rel.target_part
            if part.partname.ext.lower() in ("jpeg", "jpg"):
                part._blob = blob
                return
            package.rels.pop(rel.rId)
    part = Part(PackURI("/docProps/thumbnail.jpeg"), "image/jpeg", blob, package)
    package.relate_to(part, RT.THUMBNAIL)


def add_list(cell, title, lines, label):
    from docx.shared import Pt

    paragraph = cell.add_paragraph()
    label(paragraph, title)
    for line in lines:
        item = cell.add_paragraph(style="List Bullet")
        item.add_run(line).font.size = Pt(9)


def build(root, out=None, copy_to=None):
    try:
        from docx import Document
        from docx.enum.section import WD_ORIENT
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.shared import Cm, Pt
    except ImportError:
        fail("python-docx is missing, run: pip install python-docx pillow", 3)
    if importlib.util.find_spec("PIL") is None:
        fail("Pillow is missing, run: pip install python-docx pillow", 3)

    root = Path(root).resolve()
    doc = load_json(root / "items.json")
    validate(doc)
    labels = {**LABELS.get(doc.get("lang", "id"), LABELS["id"]), **doc.get("labels", {})}
    viewport = int(doc.get("viewport", 1440))
    boxes = {side: load_json(root / side / "boxes.json", required=False) for side in ("before", "after")}

    def label(paragraph, text, bold=True, size=10):
        run = paragraph.add_run(text)
        run.bold = bold
        run.font.size = Pt(size)
        return run

    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(1.5))

    document.add_heading(doc["title"], level=0)
    if doc.get("intro"):
        document.add_paragraph(doc["intro"])

    first_pair = None
    images = 0
    for index, item in enumerate(doc["items"], start=1):
        if index > 1:
            document.add_page_break()
        document.add_heading(f"{index}. {item['title']}", level=1)

        info = document.add_table(rows=1, cols=3)
        info.style = "Table Grid"
        add_list(info.cell(0, 0), labels["asked"], item.get("asked", []), label)
        add_list(info.cell(0, 1), labels["changed"], item.get("changed", []), label)
        add_list(info.cell(0, 2), labels["checked"], item.get("checked", []), label)

        document.add_paragraph()
        shots = document.add_table(rows=2, cols=2)
        shots.style = "Table Grid"
        shots.alignment = WD_TABLE_ALIGNMENT.CENTER
        label(shots.cell(0, 0).paragraphs[0], labels["before"])
        label(shots.cell(0, 1).paragraphs[0], labels["after"])
        pair = []
        for column, side in enumerate(("before", "after")):
            cell = shots.cell(1, column)
            relative = item.get(side)
            path = boxed_image(root, side, relative, boxes[side], viewport) if relative else None
            pair.append(path)
            if path:
                cell.paragraphs[0].add_run().add_picture(str(path), width=Cm(12.4))
                images += 1
            else:
                label(cell.paragraphs[0], item.get(f"{side}_note", labels["no_image"]), bold=False, size=9)
        if first_pair is None and any(pair):
            first_pair = pair

    author = display_name()
    now = datetime.datetime.now()
    props = document.core_properties
    props.author = author
    props.last_modified_by = author
    props.title = doc["title"]
    props.created = now
    props.modified = now
    props.revision = 1

    embed_thumbnail(document, thumbnail(doc, first_pair or [None, None], labels))

    output = Path(out) if out else root / doc.get("output", f"{doc['title']}.docx")
    document.save(output)
    print(f"OUTPUT: {output}")
    print(f"AUTHOR: {author}")
    print(f"ITEMS: {len(doc['items'])}")
    print(f"IMAGES: {images}")
    print("THUMBNAIL: embedded")
    if copy_to:
        target_dir = Path(os.path.expanduser(copy_to))
        if not target_dir.is_dir():
            fail(f"copy target {target_dir} is not a folder")
        copied = shutil.copy2(output, target_dir / output.name)
        print(f"COPIED: {copied}")
    return output


def check(path):
    path = Path(path)
    if not path.exists():
        fail(f"{path} not found")
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        core = archive.read("docProps/core.xml").decode("utf-8") if "docProps/core.xml" in names else ""
        rels = archive.read("_rels/.rels").decode("utf-8") if "_rels/.rels" in names else ""
    import re

    creator = re.search(r"<dc:creator>(.*?)</dc:creator>", core)
    print(f"AUTHOR: {creator.group(1) if creator else ''}")
    has_thumb = any(name.startswith("docProps/thumbnail") for name in names) and "metadata/thumbnail" in rels
    print(f"THUMBNAIL: {'embedded' if has_thumb else 'missing'}")
    print(f"IMAGES: {sum(1 for name in names if name.startswith('word/media/'))}")
    author = (creator.group(1) if creator else "").strip().lower()
    problems = []
    if not author or author == "python-docx":
        problems.append("author is empty or python-docx")
    if not has_thumb:
        problems.append("no thumbnail, File Explorer shows a blank page")
    for problem in problems:
        print(f"PROBLEM: {problem}")
    sys.exit(1 if problems else 0)


def init(root, title, lang):
    root = Path(root)
    if (root / "items.json").exists():
        fail(f"{root / 'items.json'} already exists")
    for side in ("before", "after"):
        (root / side).mkdir(parents=True, exist_ok=True)
        boxes = root / side / "boxes.json"
        if not boxes.exists():
            boxes.write_text("{}\n", encoding="utf-8")
    labels = LABELS.get(lang, LABELS["id"])
    skeleton = {
        "title": title,
        "intro": "",
        "lang": lang,
        "viewport": 1440,
        "output": f"{title}.docx",
        "items": [
            {
                "title": "",
                "asked": [],
                "changed": [],
                "checked": [],
                "before": "01-screen.png",
                "after": "01-screen.png",
            }
        ],
    }
    (root / "items.json").write_text(json.dumps(skeleton, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"CREATED: {root}")
    print(f"LABELS: {labels['asked']}, {labels['changed']}, {labels['checked']}")


def main(argv=None):
    parser = argparse.ArgumentParser(prog="revision-doc", description="Build a before and after revision report as a Word document.")
    sub = parser.add_subparsers(dest="command", required=True)
    p_init = sub.add_parser("init", help="create a revision folder with before, after and items.json")
    p_init.add_argument("folder")
    p_init.add_argument("--title", default="Revision report")
    p_init.add_argument("--lang", default="id", choices=sorted(LABELS))
    p_build = sub.add_parser("build", help="build the docx from a revision folder")
    p_build.add_argument("folder")
    p_build.add_argument("--out")
    p_build.add_argument("--copy-to")
    p_check = sub.add_parser("check", help="check author, thumbnail and images of a docx")
    p_check.add_argument("docx")
    args = parser.parse_args(argv)
    if args.command == "init":
        init(args.folder, args.title, args.lang)
    elif args.command == "build":
        build(args.folder, args.out, args.copy_to)
    else:
        check(args.docx)


if __name__ == "__main__":
    main()
